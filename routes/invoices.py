"""
Invoice Management Routes
"""
from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, session, make_response, send_file, current_app
from flask_login import login_required, current_user
from models import Invoice, Client, Agreement, InvoiceLineItem
from extensions import db
from datetime import datetime, date, timedelta
from decimal import Decimal
from functools import wraps
import logging
from services.secure_upload import sanitize_upload_filename
import os
import json

logger = logging.getLogger(__name__)
invoices = Blueprint('invoices', __name__)

def handle_errors(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        try:
            return f(*args, **kwargs)
        except Exception as e:
            logger.error(f"Error in {f.__name__}: {str(e)}", exc_info=True)
            flash(f'An error occurred: {str(e)}', 'error')
            return redirect(url_for('invoices.list_invoices'))
    return decorated_function


def _require_manager_or_admin():
    if not current_user.is_authenticated:
        return redirect(url_for("auth.login"))
    if not (getattr(current_user, "is_admin", False) or getattr(current_user, "is_manager", False)):
        flash("Access denied. Manager or Admin privileges required.", "error")
        return redirect(url_for("main.dashboard"))
    return None


def require_manager_or_admin(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        resp = _require_manager_or_admin()
        if resp is not None:
            return resp
        return f(*args, **kwargs)

    return decorated


@invoices.route("/upload", methods=["GET", "POST"])
@login_required
@require_manager_or_admin
@handle_errors
def upload_invoices():
    """Step 1: upload file → preview with client matching (confirm on next screen)."""
    from access_control import get_accessible_clients
    from services.invoice_import_service import build_invoice_import_preview

    if request.method == "POST":
        if "file" not in request.files:
            flash("No file selected.", "error")
            return redirect(url_for("invoices.upload_invoices"))
        f = request.files["file"]
        if not f.filename:
            flash("No file selected.", "error")
            return redirect(url_for("invoices.upload_invoices"))
        try:
            safe_name = sanitize_upload_filename(f.filename)
        except ValueError as exc:
            flash(str(exc), "error")
            return redirect(url_for("invoices.upload_invoices"))
        low = safe_name.lower()
        if not (
            low.endswith(".csv")
            or low.endswith(".tsv")
            or low.endswith(".txt")
            or low.endswith(".xls")
            or low.endswith(".xlsx")
        ):
            flash("Please upload a .csv, .tsv, .xls, or .xlsx file.", "error")
            return redirect(url_for("invoices.upload_invoices"))
        raw = f.read()
        if not raw:
            flash("File is empty.", "error")
            return redirect(url_for("invoices.upload_invoices"))

        accessible = get_accessible_clients()
        allowed_ids = {c.id for c in accessible}

        result = build_invoice_import_preview(raw, safe_name, allowed_ids)
        if not result.get("ok"):
            flash(result.get("error") or "Could not read file.", "error")
            return redirect(url_for("invoices.upload_invoices"))

        session["invoice_import_preview_id"] = result["preview_id"]
        flash(
            f"Parsed file: {result.get('row_count', 0)} row(s) ready to confirm, "
            f"{result.get('skipped_duplicate', 0)} skipped as duplicate(s) (already in database or repeated in this file), "
            f"{result.get('skipped_empty', 0)} empty invoice number line(s). "
            "Match each file name to a database client, then confirm. "
            "You can download skipped/invalid rows as CSV on the confirm screen.",
            "info",
        )
        return redirect(url_for("invoices.confirm_invoice_upload"))

    return render_template("invoices/upload.html")


@invoices.route("/upload/confirm", methods=["GET", "POST"])
@login_required
@require_manager_or_admin
@handle_errors
def confirm_invoice_upload():
    """Step 2: show file customer vs DB client; user picks client and confirms import."""
    from access_control import get_accessible_clients
    from services.invoice_import_service import (
        load_invoice_import_preview,
        delete_invoice_import_preview,
        commit_invoice_import_with_mappings,
        failure_commit_rows_for_csv,
        save_invoice_import_failure_export,
    )

    preview_id = session.get("invoice_import_preview_id")
    if not preview_id:
        flash("No import preview found. Upload a file first.", "warning")
        return redirect(url_for("invoices.upload_invoices"))

    data = load_invoice_import_preview(preview_id)
    if not data:
        session.pop("invoice_import_preview_id", None)
        flash("Preview expired or missing. Upload again.", "warning")
        return redirect(url_for("invoices.upload_invoices"))

    accessible_clients = get_accessible_clients()

    if request.method == "POST":
        posted_pid = (request.form.get("preview_id") or "").strip()
        if posted_pid != preview_id:
            flash("Preview mismatch. Start over.", "error")
            return redirect(url_for("invoices.upload_invoices"))

        include_set = set()
        client_map: dict = {}
        for key in request.form:
            if key.startswith("include_") and request.form.get(key) == "on":
                try:
                    pidx = int(key.replace("include_", ""))
                    include_set.add(pidx)
                except ValueError:
                    continue

        for pr in data.get("rows") or []:
            if pr.get("status") != "ready":
                continue
            pidx = pr["preview_index"]
            raw_cid = request.form.get(f"client_id_{pidx}")
            if raw_cid and raw_cid.strip().isdigit():
                cid = int(raw_cid)
                allowed = {c.id for c in accessible_clients}
                if cid in allowed:
                    client_map[pidx] = cid

        result = commit_invoice_import_with_mappings(
            data["rows"], client_map, include_set, current_user.id
        )
        fail_export_id = None
        if result.get("errors"):
            fail_csv_rows = failure_commit_rows_for_csv(result["errors"])
            fail_export_id = save_invoice_import_failure_export(fail_csv_rows)
            if fail_export_id:
                session["invoice_import_failed_export_id"] = fail_export_id
        else:
            session.pop("invoice_import_failed_export_id", None)

        delete_invoice_import_preview(preview_id)
        session.pop("invoice_import_preview_id", None)

        missing = result.get("clients_missing_agreement") or []
        if missing:
            session["invoice_import_missing_agreements"] = missing

        if result["created"] and not result["errors"]:
            msg_cat = "success"
        elif result["created"]:
            msg_cat = "warning"
        else:
            msg_cat = "warning"

        flash(
            f"Import complete: {result['created']} invoice(s) created, "
            f"{len(result['errors'])} error(s).",
            msg_cat,
        )
        sk_dup = result.get("skipped_duplicate_at_commit") or []
        if sk_dup:
            sample = ", ".join(
                f"row {x.get('row')} ({x.get('invoice_number', '')})" for x in sk_dup[:8]
            )
            more = f" (+{len(sk_dup) - 8} more)" if len(sk_dup) > 8 else ""
            flash(
                f"{len(sk_dup)} row(s) skipped — invoice number already in the database "
                f"(created after preview or duplicate in the same run): {sample}{more}.",
                "info",
            )
        for err in result["errors"][:20]:
            flash(f"Row {err['row']}: {err['message']}", "error")
        if len(result["errors"]) > 20:
            flash(f"... and {len(result['errors']) - 20} more errors.", "error")
        if fail_export_id:
            flash(
                "Download a CSV of rows that failed at commit from the invoice list "
                '(banner: "Download failed rows").',
                "info",
            )
        return redirect(url_for("invoices.list_invoices"))

    ready_rows = [r for r in data.get("rows") or [] if r.get("status") == "ready"]
    other_rows = [r for r in data.get("rows") or [] if r.get("status") != "ready"]

    skipped_export_count = sum(
        1 for r in other_rows if r.get("status") in ("skip_duplicate", "error")
    )

    return render_template(
        "invoices/upload_confirm.html",
        preview_id=preview_id,
        filename=data.get("filename") or "",
        ready_rows=ready_rows,
        other_rows=other_rows,
        skipped_export_count=skipped_export_count,
        clients=accessible_clients,
    )


@invoices.route("/upload/preview/export-skipped")
@login_required
@require_manager_or_admin
@handle_errors
def export_invoice_import_skipped():
    """CSV of duplicate + validation-error rows from the current preview (for fix / re-upload)."""
    from services.invoice_import_service import (
        load_invoice_import_preview,
        build_skipped_rows_for_csv,
        invoice_export_rows_to_csv_bytes,
    )

    preview_id = (request.args.get("preview_id") or "").strip()
    if not preview_id or preview_id != session.get("invoice_import_preview_id"):
        flash("Invalid or expired preview.", "error")
        return redirect(url_for("invoices.upload_invoices"))
    data = load_invoice_import_preview(preview_id)
    if not data:
        flash("Preview not found.", "error")
        return redirect(url_for("invoices.upload_invoices"))
    rows = build_skipped_rows_for_csv(data.get("rows") or [])
    if not rows:
        flash("No skipped or invalid rows to export.", "info")
        return redirect(url_for("invoices.confirm_invoice_upload"))
    blob = invoice_export_rows_to_csv_bytes(rows)
    from services.audit_service import log_data_export

    log_data_export(
        "invoice_import_skipped_csv",
        details={"preview_id": preview_id, "row_count": len(rows)},
    )
    base = (data.get("filename") or "import").rsplit(".", 1)[0]
    safe_base = "".join(c for c in base if c.isalnum() or c in "._- ")[:80] or "import"
    fn = f"{safe_base}_skipped_rows.csv"
    resp = make_response(blob)
    resp.headers["Content-Type"] = "text/csv; charset=utf-8"
    resp.headers["Content-Disposition"] = f'attachment; filename="{fn}"'
    return resp


@invoices.route("/upload/download-failed-import")
@login_required
@require_manager_or_admin
@handle_errors
def download_invoice_import_failures():
    """One-time CSV of rows that failed during commit (after confirm)."""
    from services.invoice_import_service import (
        consume_invoice_import_failure_export,
        invoice_export_rows_to_csv_bytes,
    )

    fid = session.get("invoice_import_failed_export_id")
    if not fid:
        flash("No failed-import CSV is available.", "warning")
        return redirect(url_for("invoices.list_invoices"))
    rows = consume_invoice_import_failure_export(fid)
    session.pop("invoice_import_failed_export_id", None)
    if not rows:
        flash("Failed-import file expired or was already downloaded.", "warning")
        return redirect(url_for("invoices.list_invoices"))
    blob = invoice_export_rows_to_csv_bytes(rows)
    from services.audit_service import log_data_export

    log_data_export(
        "invoice_import_failures_csv",
        details={"export_id": fid, "row_count": len(rows)},
    )
    resp = make_response(blob)
    resp.headers["Content-Type"] = "text/csv; charset=utf-8"
    resp.headers["Content-Disposition"] = (
        'attachment; filename="invoice_import_commit_failures.csv"'
    )
    return resp


@invoices.route("/upload/cancel", methods=["POST"])
@login_required
@require_manager_or_admin
@handle_errors
def cancel_invoice_upload():
    from services.invoice_import_service import delete_invoice_import_preview

    preview_id = session.pop("invoice_import_preview_id", None)
    if preview_id:
        delete_invoice_import_preview(preview_id)
    flash("Import cancelled.", "info")
    return redirect(url_for("invoices.upload_invoices"))


@invoices.route('/')
@login_required
@require_manager_or_admin
@handle_errors
def list_invoices():
    """List all invoices with filters"""
    # Get accessible clients
    from access_control import get_accessible_clients
    accessible_clients = get_accessible_clients()
    accessible_client_ids = [c.id for c in accessible_clients]
    
    # Get filter parameters
    status_filter = request.args.get('status', 'all')
    client_id = request.args.get('client_id', type=int)
    date_from = request.args.get('date_from')
    date_to = request.args.get('date_to')
    
    # Build query
    query = Invoice.query.filter(Invoice.client_id.in_(accessible_client_ids))
    
    # Apply filters
    if status_filter != 'all':
        query = query.filter(Invoice.status == status_filter)
    
    if client_id:
        query = query.filter(Invoice.client_id == client_id)
    
    if date_from:
        try:
            date_from_obj = datetime.strptime(date_from, '%Y-%m-%d').date()
            query = query.filter(Invoice.invoice_date >= date_from_obj)
        except ValueError:
            pass
    
    if date_to:
        try:
            date_to_obj = datetime.strptime(date_to, '%Y-%m-%d').date()
            query = query.filter(Invoice.invoice_date <= date_to_obj)
        except ValueError:
            pass
    
    # Get invoices
    invoices_list = query.order_by(Invoice.invoice_date.desc()).all()
    
    # Calculate statistics
    total_invoices = Invoice.query.filter(Invoice.client_id.in_(accessible_client_ids)).count()
    pending_invoices = Invoice.query.filter(
        Invoice.client_id.in_(accessible_client_ids),
        Invoice.status.in_(['draft', 'sent'])
    ).count()
    overdue_invoices = Invoice.query.filter(
        Invoice.client_id.in_(accessible_client_ids),
        Invoice.status.in_(['sent']),
        Invoice.due_date < date.today()
    ).count()
    paid_invoices = Invoice.query.filter(
        Invoice.client_id.in_(accessible_client_ids),
        Invoice.status == 'paid'
    ).count()
    
    # Calculate total amounts
    pending_amount = db.session.query(db.func.sum(Invoice.total_amount)).filter(
        Invoice.client_id.in_(accessible_client_ids),
        Invoice.status.in_(['draft', 'sent'])
    ).scalar() or Decimal('0.0')
    
    overdue_amount = db.session.query(db.func.sum(Invoice.total_amount)).filter(
        Invoice.client_id.in_(accessible_client_ids),
        Invoice.status == 'sent',
        Invoice.due_date < date.today()
    ).scalar() or Decimal('0.0')
    
    # Get all clients for filter dropdown
    clients = accessible_clients
    
    today = date.today()

    missing_agreement_clients = session.pop("invoice_import_missing_agreements", None) or []
    show_failed_import_download = bool(session.get("invoice_import_failed_export_id"))

    return render_template('invoices/list.html',
                         invoices=invoices_list,
                         clients=clients,
                         status_filter=status_filter,
                         selected_client_id=client_id,
                         date_from=date_from,
                         date_to=date_to,
                         total_invoices=total_invoices,
                         pending_invoices=pending_invoices,
                         overdue_invoices=overdue_invoices,
                         paid_invoices=paid_invoices,
                         pending_amount=pending_amount,
                         overdue_amount=overdue_amount,
                         today=today,
                         missing_agreement_clients=missing_agreement_clients,
                         show_failed_import_download=show_failed_import_download)


@invoices.route("/generate", methods=["GET", "POST"])
@login_required
@handle_errors
def generate_invoice_from_agreement():
    """Generate an invoice using agreement terms + current portfolio valuation."""
    from access_control import get_accessible_clients
    from invoice_generation_service import InvoiceGenerationService
    from sqlalchemy import func
    from sqlalchemy.orm import joinedload
    from models import Lead

    accessible_clients = get_accessible_clients()
    allowed_ids = [c.id for c in accessible_clients]

    agreements = (
        Agreement.query.options(joinedload(Agreement.lead).joinedload(Lead.client))
        .join(Lead, Agreement.lead_id == Lead.id)
        .join(Client, Lead.client_id == Client.id)
        .filter(Lead.client_id.in_(allowed_ids))
        .order_by(func.lower(Client.name), Agreement.id.desc())
        .all()
    )

    if request.method == "POST":
        raw_agreement_id = (request.form.get("agreement_id") or "").strip()
        billing_date_raw = (request.form.get("billing_date") or "").strip()

        if not raw_agreement_id.isdigit():
            flash("Please select an agreement.", "error")
            return render_template("invoices/generate.html", agreements=agreements)

        agreement_id = int(raw_agreement_id)
        selected = Agreement.query.get(agreement_id)
        if not selected or not selected.lead or not selected.lead.client_id:
            flash("Agreement is invalid or not linked to a client.", "error")
            return render_template("invoices/generate.html", agreements=agreements)

        if selected.lead.client_id not in allowed_ids:
            flash("You do not have permission to bill this client.", "error")
            return render_template("invoices/generate.html", agreements=agreements)

        billing_date = None
        if billing_date_raw:
            try:
                billing_date = datetime.strptime(billing_date_raw, "%Y-%m-%d").date()
            except ValueError:
                flash("Invalid billing date format.", "error")
                return render_template("invoices/generate.html", agreements=agreements)

        result = InvoiceGenerationService.generate_invoice(
            agreement_id=agreement_id,
            billing_date=billing_date,
            user_id=current_user.id,
        )
        if "error" in result:
            flash(str(result["error"]), "error")
            return render_template("invoices/generate.html", agreements=agreements)

        flash("Invoice generated successfully (draft).", "success")
        return redirect(url_for("invoices.view_invoice", invoice_id=result["invoice_id"]))

    return render_template("invoices/generate.html", agreements=agreements)


@invoices.route('/manager-plan', methods=['GET', 'POST'])
@login_required
@require_manager_or_admin
@handle_errors
def manager_plan():
    """On-request billing manager: weekly prioritization + catch-up roadmap."""
    from access_control import get_accessible_clients
    from services.billing_manager_service import BillingManagerService

    accessible_clients = get_accessible_clients()
    allowed_ids = [c.id for c in accessible_clients]
    plan = None

    params = {
        "as_of_date": date.today().isoformat(),
        "weekly_capacity": 5,
        "value_weight": 0.7,
        "overdue_weight": 0.3,
    }

    if request.method == 'POST':
        as_of_raw = (request.form.get('as_of_date') or '').strip()
        weekly_raw = (request.form.get('weekly_capacity') or '5').strip()
        value_raw = (request.form.get('value_weight') or '0.7').strip()
        overdue_raw = (request.form.get('overdue_weight') or '0.3').strip()

        if as_of_raw:
            try:
                as_of_obj = datetime.strptime(as_of_raw, '%Y-%m-%d').date()
                params["as_of_date"] = as_of_obj.isoformat()
            except ValueError:
                flash("Invalid as-of date format; using today.", "warning")
                as_of_obj = date.today()
                params["as_of_date"] = as_of_obj.isoformat()
        else:
            as_of_obj = date.today()
            params["as_of_date"] = as_of_obj.isoformat()

        try:
            params["weekly_capacity"] = max(1, min(int(weekly_raw), 5))
        except ValueError:
            params["weekly_capacity"] = 5
            flash("Invalid weekly capacity; defaulting to 5.", "warning")

        try:
            params["value_weight"] = max(0.0, float(value_raw))
        except ValueError:
            params["value_weight"] = 0.7
            flash("Invalid value weight; defaulting to 0.7.", "warning")
        try:
            params["overdue_weight"] = max(0.0, float(overdue_raw))
        except ValueError:
            params["overdue_weight"] = 0.3
            flash("Invalid overdue weight; defaulting to 0.3.", "warning")

        plan = BillingManagerService.run_plan(
            as_of_date=as_of_obj,
            weekly_capacity=params["weekly_capacity"],
            value_weight=params["value_weight"],
            overdue_weight=params["overdue_weight"],
            accessible_client_ids=allowed_ids,
        )

    return render_template(
        'invoices/manager_plan.html',
        params=params,
        plan=plan,
    )


@invoices.route('/<int:invoice_id>')
@login_required
@handle_errors
def view_invoice(invoice_id):
    """View invoice details"""
    invoice = Invoice.query.get_or_404(invoice_id)
    
    # Check access
    from access_control import get_accessible_clients
    accessible_clients = get_accessible_clients()
    accessible_client_ids = [c.id for c in accessible_clients]
    
    if invoice.client_id not in accessible_client_ids:
        flash('You do not have permission to view this invoice.', 'error')
        return redirect(url_for('invoices.list_invoices'))
    
    # Get line items
    line_items = InvoiceLineItem.query.filter_by(invoice_id=invoice_id).all()

    zoho = {"configured": False, "linked": False, "zoho_invoice_id": None, "zoho_payment_id": None}
    try:
        from services.zoho_books_service import invoice_zoho_summary

        zoho = invoice_zoho_summary(invoice)
    except Exception:
        pass

    return render_template(
        'invoices/view.html',
        invoice=invoice,
        line_items=line_items,
        zoho=zoho,
    )


def _load_invoice_portfolio_snapshot(invoice_id: int):
    """
    Best-effort load of the as-of portfolio snapshot captured at invoice generation time.
    Stored under the protected upload root, not under static.
    """
    try:
        upload_root = current_app.config.get("UPLOAD_FOLDER") or os.path.join(
            current_app.root_path, "uploads"
        )
        path = os.path.join(upload_root, "invoices", "portfolio_snapshots", f"invoice_{int(invoice_id)}.json")
        if not os.path.isfile(path):
            return None
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        logger.warning("Could not load invoice portfolio snapshot invoice_id=%s: %s", invoice_id, e)
        return None


@invoices.route('/<int:invoice_id>/pdf')
@login_required
@handle_errors
def download_invoice_pdf(invoice_id):
    """
    Generate and download invoice PDF, including portfolio snapshot as-of valuation date.

    Prefer WeasyPrint when available. If its system deps are missing (common on macOS
    without GTK/Pango), fall back to a pure-Python ReportLab PDF.
    """
    invoice = Invoice.query.get_or_404(invoice_id)

    # Access check
    from access_control import get_accessible_clients
    accessible_clients = get_accessible_clients()
    accessible_client_ids = [c.id for c in accessible_clients]
    if invoice.client_id not in accessible_client_ids:
        flash('You do not have permission to view this invoice.', 'error')
        return redirect(url_for('invoices.list_invoices'))

    line_items = InvoiceLineItem.query.filter_by(invoice_id=invoice_id).all()
    snapshot = _load_invoice_portfolio_snapshot(invoice_id)

    # Render HTML
    html = render_template(
        "invoices/pdf.html",
        invoice=invoice,
        line_items=line_items,
        snapshot=snapshot,
    )

    pdf_bytes = None
    weasy_err = None
    try:
        from weasyprint import HTML  # type: ignore
        pdf_bytes = HTML(string=html, base_url=current_app.root_path).write_pdf()
    except Exception as e:
        weasy_err = e

    if pdf_bytes is None:
        # Fallback: ReportLab (no external shared-library deps)
        try:
            from io import BytesIO
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import cm
            from reportlab.platypus import (
                SimpleDocTemplate,
                Paragraph,
                Spacer,
                Table,
                TableStyle,
            )

            buf = BytesIO()
            doc = SimpleDocTemplate(
                buf,
                pagesize=A4,
                leftMargin=1.6 * cm,
                rightMargin=1.6 * cm,
                topMargin=1.6 * cm,
                bottomMargin=1.6 * cm,
                title=f"Invoice {invoice.invoice_number}",
            )
            styles = getSampleStyleSheet()
            story = []

            story.append(Paragraph(f"<b>Invoice</b> {invoice.invoice_number}", styles["Title"]))
            story.append(Spacer(1, 8))
            story.append(
                Paragraph(
                    f"<b>Client:</b> {invoice.client.name} &nbsp;&nbsp; "
                    f"<b>Invoice date:</b> {invoice.invoice_date.strftime('%Y-%m-%d')} &nbsp;&nbsp; "
                    f"<b>Due:</b> {invoice.due_date.strftime('%Y-%m-%d')}",
                    styles["Normal"],
                )
            )
            story.append(
                Paragraph(
                    f"<b>Billing period:</b> {invoice.billing_period_start.strftime('%Y-%m-%d')} → "
                    f"{invoice.billing_period_end.strftime('%Y-%m-%d')}",
                    styles["Normal"],
                )
            )
            story.append(Spacer(1, 10))

            # Line items table
            data = [["Asset class", "Portfolio value", "Rate", "Calculated fee", "Final fee"]]
            for it in line_items:
                data.append(
                    [
                        it.asset_class_name,
                        f"₹{float(it.portfolio_value):,.2f}",
                        f"{float(it.rate_percentage) * 100:.2f}%",
                        f"₹{float(it.calculated_fee):,.2f}",
                        f"₹{float(it.final_fee):,.2f}",
                    ]
                )
            t = Table(data, hAlign="LEFT", colWidths=[5.0 * cm, 3.2 * cm, 2.0 * cm, 3.0 * cm, 3.0 * cm])
            t.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2ff")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
                        ("ALIGN", (0, 0), (0, -1), "LEFT"),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("FONTSIZE", (0, 0), (-1, -1), 9),
                        ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
                        ("TOPPADDING", (0, 0), (-1, 0), 6),
                    ]
                )
            )
            story.append(Paragraph("<b>Line items</b>", styles["Heading3"]))
            story.append(t)
            story.append(Spacer(1, 10))

            # Totals
            totals = Table(
                [
                    ["Net amount", f"₹{float(invoice.net_amount):,.2f}"],
                    [f"Tax ({float(invoice.tax_rate) * 100:.2f}%)", f"₹{float(invoice.tax_amount):,.2f}"],
                    ["Total", f"₹{float(invoice.total_amount):,.2f}"],
                ],
                hAlign="RIGHT",
                colWidths=[4.5 * cm, 4.0 * cm],
            )
            totals.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                        ("FONTNAME", (0, 0), (-1, -2), "Helvetica"),
                        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#f1f5f9")),
                        ("FONTSIZE", (0, 0), (-1, -1), 9),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story.append(totals)

            # Portfolio snapshot section (optional)
            if snapshot:
                story.append(Spacer(1, 14))
                as_of = snapshot.get("as_of") or ""
                total_v = float(snapshot.get("total_value") or 0.0)
                story.append(Paragraph(f"<b>Portfolio valuation (as-of {as_of})</b>", styles["Heading3"]))
                story.append(Paragraph(f"<b>Total value:</b> ₹{total_v:,.2f}", styles["Normal"]))
                holds = snapshot.get("holdings") or []
                if holds:
                    snap_rows = [["Security", "Asset class", "Qty", "Value"]]
                    for h in holds[:250]:
                        sym = h.get("symbol") or h.get("security_symbol") or h.get("name") or "—"
                        acn = h.get("asset_class") or "—"
                        qty = h.get("quantity") or h.get("qty") or "—"
                        val = float(h.get("current_value") or h.get("value") or 0.0)
                        snap_rows.append([str(sym), str(acn), str(qty), f"₹{val:,.2f}"])
                    snap_table = Table(
                        snap_rows,
                        hAlign="LEFT",
                        colWidths=[6.0 * cm, 3.5 * cm, 2.0 * cm, 4.2 * cm],
                    )
                    snap_table.setStyle(
                        TableStyle(
                            [
                                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f8fafc")),
                                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e2e8f0")),
                                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                                ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
                                ("FONTSIZE", (0, 0), (-1, -1), 8),
                            ]
                        )
                    )
                    story.append(Spacer(1, 6))
                    story.append(snap_table)

            if invoice.notes:
                story.append(Spacer(1, 10))
                story.append(Paragraph("<b>Notes</b>", styles["Heading3"]))
                story.append(Paragraph(str(invoice.notes), styles["Normal"]))

            doc.build(story)
            pdf_bytes = buf.getvalue()
        except Exception as e:
            msg = f"PDF generation failed (WeasyPrint: {weasy_err}; ReportLab: {e})"
            flash(msg, "error")
            return redirect(url_for('invoices.view_invoice', invoice_id=invoice_id))

    # Persist under static/ and store path on invoice for later
    out_dir = os.path.join(current_app.root_path, "static", "invoices", "pdfs")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, f"invoice_{invoice_id}.pdf")
    with open(out_path, "wb") as f:
        f.write(pdf_bytes)

    invoice.pdf_path = f"/static/invoices/pdfs/invoice_{invoice_id}.pdf"
    db.session.commit()

    download_name = f"{invoice.invoice_number}.pdf"
    return send_file(out_path, as_attachment=True, download_name=download_name, mimetype="application/pdf")


@invoices.route('/<int:invoice_id>/portfolio.xlsx')
@login_required
@handle_errors
def download_invoice_portfolio_excel(invoice_id):
    """Download invoice portfolio snapshot as an Excel file (as-of valuation date)."""
    invoice = Invoice.query.get_or_404(invoice_id)

    # Access check
    from access_control import get_accessible_clients
    accessible_clients = get_accessible_clients()
    accessible_client_ids = [c.id for c in accessible_clients]
    if invoice.client_id not in accessible_client_ids:
        flash('You do not have permission to view this invoice.', 'error')
        return redirect(url_for('invoices.list_invoices'))

    snapshot = _load_invoice_portfolio_snapshot(invoice_id)
    if not snapshot:
        flash('No portfolio snapshot found for this invoice. Re-generate invoice to capture snapshot.', 'warning')
        return redirect(url_for('invoices.view_invoice', invoice_id=invoice_id))

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
        from openpyxl.utils import get_column_letter
    except Exception as e:
        flash(f'Excel export engine not available: {e}', 'error')
        return redirect(url_for('invoices.view_invoice', invoice_id=invoice_id))

    as_of = snapshot.get("as_of") or ""
    total_value = float(snapshot.get("total_value") or 0.0)
    holdings = snapshot.get("holdings") or []

    wb = Workbook()
    ws = wb.active
    ws.title = "Portfolio"

    # Header
    ws["A1"] = "Invoice"
    ws["B1"] = invoice.invoice_number
    ws["A2"] = "Client"
    ws["B2"] = invoice.client.name if invoice.client else ""
    ws["A3"] = "As-of date"
    ws["B3"] = as_of
    ws["A4"] = "Total value (₹)"
    ws["B4"] = total_value
    for cell in ("A1", "A2", "A3", "A4"):
        ws[cell].font = Font(bold=True)

    ws.append([])
    start_row = ws.max_row + 1

    columns = [
        ("Symbol", "symbol"),
        ("Name", "name"),
        ("Asset Class", "asset_class"),
        ("Quantity", "quantity"),
        ("Price", "price"),
        ("Value (₹)", "current_value"),
    ]
    for idx, (label, _) in enumerate(columns, start=1):
        c = ws.cell(row=start_row, column=idx, value=label)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")

    def _get(h, key):
        # tolerate differing payload keys
        if key == "symbol":
            return h.get("symbol") or h.get("security_symbol") or ""
        if key == "quantity":
            return h.get("quantity") or h.get("qty") or ""
        if key == "current_value":
            return h.get("current_value") or h.get("value") or 0
        return h.get(key) or ""

    row = start_row + 1
    for h in holdings:
        ws.cell(row=row, column=1, value=str(_get(h, "symbol")))
        ws.cell(row=row, column=2, value=str(_get(h, "name")))
        ws.cell(row=row, column=3, value=str(_get(h, "asset_class")))
        ws.cell(row=row, column=4, value=_get(h, "quantity"))
        ws.cell(row=row, column=5, value=h.get("price") or h.get("current_price") or "")
        ws.cell(row=row, column=6, value=float(_get(h, "current_value") or 0))
        row += 1

    # Column widths
    for i in range(1, 7):
        ws.column_dimensions[get_column_letter(i)].width = [18, 30, 18, 12, 12, 16][i - 1]

    # Persist under static/ so it can be re-downloaded
    out_dir = os.path.join(current_app.root_path, "static", "invoices", "portfolio_exports")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, f"invoice_{invoice_id}_portfolio_asof_{as_of or 'unknown'}.xlsx")
    wb.save(out_path)

    download_name = f"{invoice.invoice_number}_portfolio_{as_of or ''}.xlsx".replace(" ", "_")
    return send_file(
        out_path,
        as_attachment=True,
        download_name=download_name,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

@invoices.route('/<int:invoice_id>/edit', methods=['GET', 'POST'])
@login_required
@require_manager_or_admin
@handle_errors
def edit_invoice(invoice_id):
    """Edit invoice status and details"""
    invoice = Invoice.query.get_or_404(invoice_id)
    
    # Check access
    from access_control import get_accessible_clients
    accessible_clients = get_accessible_clients()
    accessible_client_ids = [c.id for c in accessible_clients]
    
    if invoice.client_id not in accessible_client_ids:
        flash('You do not have permission to edit this invoice.', 'error')
        return redirect(url_for('invoices.list_invoices'))
    
    if request.method == 'POST':
        try:
            # Update status
            new_status = request.form.get('status')
            if new_status in ['draft', 'sent', 'paid', 'overdue', 'cancelled']:
                if new_status == 'sent':
                    from services.onboarding_approval_service import invoice_may_send_to_client

                    ok, reason = invoice_may_send_to_client(invoice)
                    if not ok:
                        flash(reason, 'error')
                        return redirect(url_for('invoices.edit_invoice', invoice_id=invoice_id))
                invoice.status = new_status
            
            # Update paid date if status is paid
            if new_status == 'paid':
                paid_date_str = request.form.get('paid_date')
                if paid_date_str:
                    try:
                        invoice.paid_date = datetime.strptime(paid_date_str, '%Y-%m-%d').date()
                    except ValueError:
                        pass
                elif not invoice.paid_date:
                    invoice.paid_date = date.today()
            else:
                invoice.paid_date = None
            
            # Update payment details
            invoice.payment_reference = request.form.get('payment_reference', '')
            invoice.payment_method = request.form.get('payment_method', '')
            invoice.notes = request.form.get('notes', '')
            
            db.session.commit()
            flash('Invoice updated successfully.', 'success')
            return redirect(url_for('invoices.view_invoice', invoice_id=invoice_id))
            
        except Exception as e:
            db.session.rollback()
            logger.error(f"Error updating invoice: {str(e)}")
            flash(f'Error updating invoice: {str(e)}', 'error')
    
    return render_template('invoices/edit.html', invoice=invoice)

@invoices.route('/<int:invoice_id>/update-status', methods=['POST'])
@login_required
@require_manager_or_admin
@handle_errors
def update_invoice_status(invoice_id):
    """Update invoice status via AJAX"""
    invoice = Invoice.query.get_or_404(invoice_id)
    
    # Check access
    from access_control import get_accessible_clients
    accessible_clients = get_accessible_clients()
    accessible_client_ids = [c.id for c in accessible_clients]
    
    if invoice.client_id not in accessible_client_ids:
        return jsonify({'success': False, 'error': 'Permission denied'}), 403
    
    try:
        data = request.get_json()
        new_status = data.get('status')
        
        if new_status not in ['draft', 'sent', 'paid', 'overdue', 'cancelled']:
            return jsonify({'success': False, 'error': 'Invalid status'}), 400

        if new_status == 'sent':
            from services.onboarding_approval_service import invoice_may_send_to_client

            ok, reason = invoice_may_send_to_client(invoice)
            if not ok:
                return jsonify({'success': False, 'error': reason}), 400
        
        invoice.status = new_status
        
        if new_status == 'paid':
            if not invoice.paid_date:
                invoice.paid_date = date.today()
        else:
            invoice.paid_date = None
        
        db.session.commit()
        return jsonify({'success': True, 'message': 'Status updated successfully'})
        
    except Exception as e:
        db.session.rollback()
        logger.error(f"Error updating invoice status: {str(e)}")
        return jsonify({'success': False, 'error': str(e)}), 500


def _invoice_access_ok(invoice) -> bool:
    from access_control import get_accessible_clients

    accessible_client_ids = [c.id for c in get_accessible_clients()]
    return invoice.client_id in accessible_client_ids


@invoices.route('/<int:invoice_id>/zoho/push', methods=['POST'])
@login_required
@require_manager_or_admin
@handle_errors
def zoho_push_invoice(invoice_id):
    invoice = Invoice.query.get_or_404(invoice_id)
    if not _invoice_access_ok(invoice):
        flash('You do not have permission to update this invoice.', 'error')
        return redirect(url_for('invoices.list_invoices'))
    from services.zoho_books_service import push_invoice_to_zoho, zoho_configured

    if not zoho_configured():
        flash('Zoho is not configured.', 'warning')
        return redirect(url_for('invoices.view_invoice', invoice_id=invoice_id))
    result = push_invoice_to_zoho(invoice, send_email=True)
    if result.get('error'):
        flash(result['error'], 'error')
    else:
        msg = f"Zoho invoice {'linked' if result.get('already_linked') else 'created'} ({result.get('zoho_invoice_id')})."
        if result.get('emailed'):
            msg += ' Email sent.'
        flash(msg, 'success')
    return redirect(url_for('invoices.view_invoice', invoice_id=invoice_id))


@invoices.route('/<int:invoice_id>/zoho/sync', methods=['POST'])
@login_required
@require_manager_or_admin
@handle_errors
def zoho_sync_invoice(invoice_id):
    invoice = Invoice.query.get_or_404(invoice_id)
    if not _invoice_access_ok(invoice):
        flash('You do not have permission to view this invoice.', 'error')
        return redirect(url_for('invoices.list_invoices'))
    from services.zoho_books_service import sync_invoice_status_from_zoho, zoho_configured

    if not zoho_configured():
        flash('Zoho is not configured.', 'warning')
        return redirect(url_for('invoices.view_invoice', invoice_id=invoice_id))
    result = sync_invoice_status_from_zoho(invoice, apply_local=True)
    if result.get('error'):
        flash(result['error'], 'error')
    else:
        flash(
            f"Zoho: {result.get('zoho_status')} → local {result.get('local_status')} "
            f"(balance {result.get('balance')}).",
            'success',
        )
    return redirect(url_for('invoices.view_invoice', invoice_id=invoice_id))


@invoices.route('/<int:invoice_id>/zoho/mark-paid', methods=['POST'])
@login_required
@require_manager_or_admin
@handle_errors
def zoho_mark_paid(invoice_id):
    invoice = Invoice.query.get_or_404(invoice_id)
    if not _invoice_access_ok(invoice):
        flash('You do not have permission to update this invoice.', 'error')
        return redirect(url_for('invoices.list_invoices'))
    from services.zoho_books_service import record_payment_in_zoho, zoho_configured

    if not zoho_configured():
        flash('Zoho is not configured.', 'warning')
        return redirect(url_for('invoices.view_invoice', invoice_id=invoice_id))
    ref = (request.form.get('payment_reference') or '').strip()
    method = (request.form.get('payment_method') or '').strip()
    result = record_payment_in_zoho(
        invoice,
        payment_reference=ref,
        payment_method=method,
        mark_local_paid=True,
    )
    if result.get('error'):
        flash(result['error'], 'error')
    else:
        flash(
            f"Payment recorded in Zoho (payment {result.get('zoho_payment_id')}). "
            f"Local status: paid.",
            'success',
        )
    return redirect(url_for('invoices.view_invoice', invoice_id=invoice_id))


@invoices.route('/<int:invoice_id>/zoho/pdf')
@login_required
@handle_errors
def zoho_invoice_pdf(invoice_id):
    invoice = Invoice.query.get_or_404(invoice_id)
    if not _invoice_access_ok(invoice):
        flash('You do not have permission to view this invoice.', 'error')
        return redirect(url_for('invoices.list_invoices'))
    from flask import Response
    from services.zoho_books_service import download_zoho_invoice_pdf, zoho_configured

    if not zoho_configured():
        flash('Zoho is not configured.', 'warning')
        return redirect(url_for('invoices.view_invoice', invoice_id=invoice_id))
    pdf, filename, err = download_zoho_invoice_pdf(invoice)
    if err:
        flash(err, 'error')
        return redirect(url_for('invoices.view_invoice', invoice_id=invoice_id))
    return Response(
        pdf,
        mimetype='application/pdf',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'},
    )


@invoices.route('/<int:invoice_id>/zoho/receipt')
@login_required
@handle_errors
def zoho_payment_receipt(invoice_id):
    invoice = Invoice.query.get_or_404(invoice_id)
    if not _invoice_access_ok(invoice):
        flash('You do not have permission to view this invoice.', 'error')
        return redirect(url_for('invoices.list_invoices'))
    from flask import Response
    from services.zoho_books_service import download_zoho_payment_receipt_pdf, zoho_configured

    if not zoho_configured():
        flash('Zoho is not configured.', 'warning')
        return redirect(url_for('invoices.view_invoice', invoice_id=invoice_id))
    pdf, filename, err = download_zoho_payment_receipt_pdf(invoice)
    if err:
        flash(err, 'error')
        return redirect(url_for('invoices.view_invoice', invoice_id=invoice_id))
    return Response(
        pdf,
        mimetype='application/pdf',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'},
    )

