"""Tests for agreement fee schedule formatting and billing config parsing."""

from decimal import Decimal

import pytest

from services.agreement_billing_presets import preset_slabs
from services.agreement_billing_config_service import (
    STRUCTURE_UNIFORM,
    infer_structure_type,
    parse_form_to_rates,
    validate_rates,
)
from services.agreement_fee_schedule_formatter import (
    format_asset_class_schedule,
    format_fee_schedule_from_rates,
)


class _FakeAssetClass:
    def __init__(self, name):
        self.name = name


class _FakeRate:
    def __init__(self, ac_name, min_amt, max_amt, rate_pct, is_active=True):
        self.asset_class = _FakeAssetClass(ac_name)
        self.asset_class_id = 1
        self.min_amount = Decimal(str(min_amt))
        self.max_amount = Decimal(str(max_amt)) if max_amt is not None else None
        self.rate_percentage = Decimal(str(rate_pct)) / Decimal("100")
        self.min_fee = Decimal("0")
        self.max_fee = None
        self.is_active = is_active


def test_standard_equity_preset_has_three_bands():
    slabs = preset_slabs("standard_equity")
    assert len(slabs) == 3
    assert float(slabs[0]["rate_pct"]) == 2.0
    assert float(slabs[2]["rate_pct"]) == 1.0


def test_parse_uniform_expands_to_all_asset_classes():
    form = {
        "billing_structure_type": STRUCTURE_UNIFORM,
        "uniform_rate_pct": "1.5",
    }
    rates = parse_form_to_rates(form, [1, 2, 3])
    assert len(rates) == 3
    for ac_id in (1, 2, 3):
        assert len(rates[ac_id]) == 1
        assert float(rates[ac_id][0]["rate_percentage"]) == 0.015


def test_parse_differential_exceptions():
    form = {
        "billing_structure_type": "differential",
        "default_rate_pct": "1.25",
        "billing_exception_count": "1",
        "billing_exception_0_ac_id": "2",
        "billing_exception_0_rate_pct": "1.0",
    }
    rates = parse_form_to_rates(form, [1, 2])
    assert float(rates[1][0]["rate_percentage"]) == 0.0125
    assert float(rates[2][0]["rate_percentage"]) == 0.01


def test_infer_uniform_from_matching_slabs():
    slabs = {
        1: [{"min_amount": 0, "max_amount": None, "rate_pct": 1.25}],
        2: [{"min_amount": 0, "max_amount": None, "rate_pct": 1.25}],
    }
    assert infer_structure_type(slabs) == STRUCTURE_UNIFORM


def test_format_equity_tiers_prose():
    rates = [
        _FakeRate("Equity", 0, 2_500_000, 2.0),
        _FakeRate("Equity", 2_500_000, 5_000_000, 1.5),
        _FakeRate("Equity", 5_000_000, None, 1.0),
    ]
    text = format_fee_schedule_from_rates(rates, billing_frequency="yearly")
    assert "Equity" in text
    assert "2%" in text
    assert "1.5%" in text
    assert "annual" in text.lower()


def test_format_groups_identical_flat_classes():
    rates = [
        _FakeRate("Debt", 0, None, 1.25),
        _FakeRate("Mutual Funds", 0, None, 1.25),
    ]
    text = format_fee_schedule_from_rates(rates)
    assert "1.25%" in text
    assert "and" in text or "," in text


def test_format_asset_class_flat_line():
    line = format_asset_class_schedule(
        "Debt",
        [{"min_amount": 0, "max_amount": None, "rate_pct": 1.25, "min_fee": 0, "max_fee": None}],
    )
    assert "1.25%" in line
    assert "Debt" in line


def test_validate_warns_empty_slabs():
    warnings = validate_rates({5: []})
    assert any("no billing slabs" in w for w in warnings)


def test_legacy_slots_differential_equity_exception():
    import json

    from services.agreement_fee_schedule_formatter import legacy_template_slot_values

    class AC:
        def __init__(self, name):
            self.name = name

    class Row:
        def __init__(self, name, rate_pct, ac_id=1):
            self.asset_class = AC(name)
            self.asset_class_id = ac_id
            self.min_amount = 0
            self.max_amount = None
            self.rate_percentage = rate_pct / 100.0
            self.min_fee = 0
            self.max_fee = None
            self.is_active = True

    class Ag:
        agreement_data = json.dumps(
            {
                "billing_structure_type": "differential",
                "default_rate_pct": "0.5",
                "billing_exceptions": [{"ac_id": 1, "rate_pct": 1.0}],
            }
        )
        billing_rates = [
            Row("Equity", 1.0, ac_id=1),
            Row("Debt", 0.5, ac_id=2),
            Row("Mutual Funds", 0.5, ac_id=3),
        ]

    slots = legacy_template_slot_values(Ag())
    assert slots["rate_1"] == "1"
    assert "equity portfolio value" in slots["amount_range1"].lower()
    assert slots["rate_2"] == "0.5"
    assert "portfolio value of" in slots["amount_range2"].lower()
    assert "rate_3" not in slots


def test_legacy_rate_tier_lines_omits_empty_third_tier():
    import json

    from services.agreement_fee_schedule_formatter import (
        format_legacy_rate_tier_lines,
        legacy_template_slot_values,
    )

    class AC:
        def __init__(self, name):
            self.name = name

    class Row:
        def __init__(self, name, rate_pct, ac_id=1):
            self.asset_class = AC(name)
            self.asset_class_id = ac_id
            self.min_amount = 0
            self.max_amount = None
            self.rate_percentage = rate_pct / 100.0
            self.min_fee = 0
            self.max_fee = None
            self.is_active = True

    class Ag:
        agreement_data = json.dumps(
            {
                "billing_structure_type": "differential",
                "default_rate_pct": "0.5",
                "billing_exceptions": [{"ac_id": 1, "rate_pct": 1.0}],
            }
        )
        billing_rates = [Row("Equity", 1.0, ac_id=1), Row("Debt", 0.5, ac_id=2)]
        variables = []

    lines = format_legacy_rate_tier_lines(Ag())
    assert "rate_3" not in legacy_template_slot_values(Ag())
    assert "Assets" in lines and "Rate" in lines
    assert "Equity" in lines and "1%" in lines
    assert "Debt" in lines and "0.5%" in lines
    assert "Fee will be calculated as" not in lines
    assert "basis" not in lines.lower()


def test_portfolio_value_fee_schedule_wording():
    import json

    from services.agreement_fee_schedule_formatter import (
        fee_schedule_asset_rate_rows,
        format_portfolio_value_fee_schedule,
    )

    class AC:
        def __init__(self, name):
            self.name = name

    class Row:
        def __init__(self, name, rate_pct, ac_id=1):
            self.asset_class = AC(name)
            self.asset_class_id = ac_id
            self.min_amount = 0
            self.max_amount = None
            self.rate_percentage = rate_pct / 100.0
            self.min_fee = 0
            self.max_fee = None
            self.is_active = True

    class Var:
        def __init__(self, name, value):
            self.variable_name = name
            self.variable_value = value
            self.variable_type = "billing_config"

    class Ag:
        agreement_data = json.dumps({"billing_structure_type": "differential"})
        billing_rates = [
            Row("Equity", 1.0, ac_id=1),
            Row("Debt ETF", 0.5, ac_id=2),
            Row("Debt Mutual Funds", 0.5, ac_id=3),
            Row("Gold ETF", 0.5, ac_id=4),
            Row("Fixed income / Debt", 0.0, ac_id=5),
        ]
        variables = [
            Var("billing_frequency", "half_yearly"),
            Var("period_start_month", "6"),
            Var("valuation_date_rule", "prepaid"),
        ]

    text = format_portfolio_value_fee_schedule(Ag())
    assert "Advisory fee structure" in text
    assert "Assets" in text and "Rate" in text
    assert "Equity" in text and "1%" in text
    assert "Fixed income / Debt" in text and "0%" in text
    assert "Debt ETF" in text and "0.5%" in text
    # Same-rate assets are merged into one Assets cell
    assert "Debt ETF and Debt Mutual Funds" in text or (
        "Debt ETF" in text and "Debt Mutual Funds" in text and text.count("0.5%") == 1
    )
    assert "Fee will be calculated as" not in text
    rows = fee_schedule_asset_rate_rows(Ag())
    assert any(r[1] == "1%" and "Equity" in r[0] for r in rows)
    assert any(r[1] == "0%" and "Fixed income / Debt" in r[0] for r in rows)
    # One row per distinct rate
    assert len(rows) == 3


def test_billing_fee_schedule_becomes_word_table(tmp_path):
    docx = pytest.importorskip("docx")
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    from services.agreement_docx_service import replace_placeholders_in_document
    from services.agreement_fee_schedule_formatter import format_fee_schedule_table_text

    doc = Document()
    p = doc.add_paragraph("<<billing_fee_schedule>>")
    # Simulate template yellow highlight on the placeholder
    from docx.enum.text import WD_COLOR_INDEX

    p.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
    rows = [
        ("Equity, Gold ETF / Silver ETF, and REIT / InvIT", "1%"),
        ("Equity ETF and Equity Mutual Funds", "0.2%"),
        ("Debt ETF, Debt Mutual Funds, and Fixed income / Debt", "0%"),
    ]
    replace_placeholders_in_document(
        doc,
        {"billing_fee_schedule": format_fee_schedule_table_text(rows)},
    )
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert table.rows[0].cells[0].text == "Assets"
    assert table.rows[0].cells[1].text == "Rate"
    assert len(table.rows) == 4  # header + 3 merged rate groups
    for row in table.rows:
        for cell in row.cells:
            assert cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
            borders = cell._tc.tcPr.find(qn("w:tcBorders"))
            assert borders is not None
            assert borders.find(qn("w:top")) is not None
    # Placeholder highlight should be cleared on the intro paragraph
    for run in doc.paragraphs[0].runs:
        assert run.font.highlight_color is None
    path = tmp_path / "fee.docx"
    doc.save(path)
    from services.agreement_docx_service import filled_docx_to_html

    html = filled_docx_to_html(str(path))
    assert "<table" in html
    assert "Fixed income / Debt" in html
    assert "Equity, Gold ETF / Silver ETF, and REIT / InvIT" in html

