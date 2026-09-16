"""Tests for agreement DOCX placeholder extraction and substitution."""
import os
import tempfile

import pytest

pytestmark = pytest.mark.no_app

from services.agreement_docx_service import (
    docx_full_text,
    extract_variables_from_text,
    replace_placeholders_in_document,
    substitute_placeholders,
)


def test_apply_inertia_logo_watermark(tmp_path):
    docx = pytest.importorskip("docx")
    pytest.importorskip("PIL")
    from docx import Document

    from services.agreement_docx_service import (
        apply_inertia_logo_watermark,
        resolve_inertia_logo_path,
    )

    logo = resolve_inertia_logo_path()
    if not logo:
        pytest.skip("inertia-logo.png not present")

    doc = Document()
    doc.add_paragraph("Agreement body <<client>>")
    assert apply_inertia_logo_watermark(doc, logo_path=logo, opacity=0.2) is True
    header_xml = doc.sections[0].header._element.xml
    assert "drawing" in header_xml or "blip" in header_xml or "pict" in header_xml
    out = tmp_path / "wm.docx"
    doc.save(out)
    assert out.stat().st_size > 1000


def test_resolve_inertia_logo_path_finds_png():
    from services.agreement_docx_service import resolve_inertia_logo_path

    path = resolve_inertia_logo_path()
    if path is None:
        pytest.skip("inertia-logo.png not present")
    assert path.endswith("inertia-logo.png")
    assert os.path.isfile(path)


def test_analyze_uploaded_template_reads_docx_placeholders(tmp_path):
    docx = pytest.importorskip("docx")
    from docx import Document
    from services.agreement_docx_service import analyze_uploaded_template

    doc = Document()
    doc.add_paragraph("Client: <<client_name>> Date: <<agreement_date>>")
    path = tmp_path / "Inertia_agreement_template.docx"
    doc.save(path)
    names = analyze_uploaded_template(str(path), path.name)
    assert names == ["client_name", "agreement_date"]


def test_analyze_uploaded_template_rejects_garbage_docx(tmp_path):
    from services.agreement_docx_service import analyze_uploaded_template

    path = tmp_path / "not-a-word.docx"
    path.write_bytes(b"not a zip")
    with pytest.raises(ValueError, match="Could not read this Word file"):
        analyze_uploaded_template(str(path), path.name)


def test_extract_variables_strips_whitespace():
    text = "Hello << client >> and <<date>>"
    assert extract_variables_from_text(text) == ['client', 'date']


def test_substitute_placeholders_allows_whitespace():
    text = "Signed: << client >> on << date >>"
    out = substitute_placeholders(text, {'client': 'Aniket Misar', 'date': '03 June 2026'})
    assert 'Aniket Misar' in out
    assert '03 June 2026' in out
    assert '<<' not in out


def test_replace_clears_placeholder_highlight():
    docx = pytest.importorskip("docx")
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    from services.agreement_docx_service import replace_placeholders_in_document

    doc = Document()
    p = doc.add_paragraph("<<advisory_fee_mode>>")
    p.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
    replace_placeholders_in_document(
        doc,
        {
            "advisory_fee_mode": (
                "3.1 Assets under Advice (AUA) mode:\n"
                "Assets Under Advice (AUA): Market value of Equity assets under advice "
                "covered by this arrangement."
            )
        },
    )
    assert "Assets Under Advice (AUA)" in doc.paragraphs[0].text
    assert "<<" not in doc.paragraphs[0].text
    for run in doc.paragraphs[0].runs:
        assert run.font.highlight_color is None


def test_replace_in_table_cell():
    docx = pytest.importorskip('docx')
    from docx import Document

    doc = Document()
    doc.add_paragraph('Body <<client>>')
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = 'Signature block: <<client>>'

    n = replace_placeholders_in_document(doc, {'client': 'PAN Legal Name'})
    assert n >= 2
    assert 'PAN Legal Name' in docx_full_text(doc)
    assert '<<client>>' not in docx_full_text(doc)


def test_extract_variables_from_docx_path_roundtrip():
    docx = pytest.importorskip('docx')
    from docx import Document

    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, 't.docx')
        doc = Document()
        doc.add_paragraph('Date: <<date>>')
        t = doc.add_table(rows=1, cols=1)
        t.rows[0].cells[0].text = '<<pan>>'
        doc.save(path)

        from services.agreement_docx_service import extract_variables_from_docx_path

        assert set(extract_variables_from_docx_path(path)) == {'date', 'pan'}


def test_filled_docx_to_html_includes_table_text():
    docx = pytest.importorskip('docx')
    from docx import Document

    from services.agreement_docx_service import filled_docx_to_html

    doc = Document()
    t = doc.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = 'Table <<client>>'
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
        path = tmp.name
    try:
        doc.save(path)
        html = filled_docx_to_html(path)
        assert 'Table' in html
    finally:
        os.unlink(path)
