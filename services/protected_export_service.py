"""
Helpers for sensitive file downloads: PII masking pass + no-cache response headers.

Export audit rows: routes call log_data_export() — gated when DEFER_DB_FEATURES includes
audit_log (see docs/DB_CUTOVER_REGISTRY.md).
"""
from __future__ import annotations

from io import BytesIO
from flask import Response, send_file

from services.pii_masking import mask_client_name, mask_email

PII_MASK_EXPORT_VERSION = "v2"


def mask_practice_analytics_workbook_bytes(data: bytes) -> bytes:
    """Ensure Clients sheet name (col B) and email (col C) are masked."""
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data))
    if "Clients" not in wb.sheetnames:
        return data
    ws = wb["Clients"]
    for row in range(2, ws.max_row + 1):
        name_val = ws.cell(row, 2).value
        email_val = ws.cell(row, 3).value
        if name_val is not None and str(name_val).strip():
            ws.cell(row, 2).value = mask_client_name(str(name_val))
        if email_val is not None and str(email_val).strip():
            ws.cell(row, 3).value = mask_email(str(email_val))
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def attachment_response(
    data: bytes,
    *,
    download_name: str,
    mimetype: str,
) -> Response:
    """File download with cache disabled (avoid stale unmasked exports in browser)."""
    resp = send_file(
        BytesIO(data),
        mimetype=mimetype,
        as_attachment=True,
        download_name=download_name,
    )
    resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    resp.headers["Pragma"] = "no-cache"
    resp.headers["X-Export-PII-Masked"] = PII_MASK_EXPORT_VERSION
    return resp
